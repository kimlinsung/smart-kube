from __future__ import annotations

import json
import os
import tempfile
import unittest
import zipfile

from backend import file_preview


class FilePreviewTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def preview(self, name, content=None):
        path = os.path.join(self.temp.name, name)
        if content is not None:
            with open(path, "wb") as handle:
                handle.write(content.encode("utf-8") if isinstance(content, str) else content)
        return file_preview.read_preview({"original_name": name, "stored_path": path})

    def test_csv_uses_csv_parser_and_bounds_rows(self):
        data = self.preview("data.csv", 'name,value\n"first, second","line 1\nline 2"\n' + 'x,2\n' * 210)
        self.assertEqual(data["sheets"][0]["rows"][1], ["first, second", "line 1\nline 2"])
        self.assertEqual(len(data["sheets"][0]["rows"]), 200)
        self.assertTrue(data["truncated"])

    def test_large_text_is_truncated_not_rejected(self):
        data = self.preview("output.log", "a" * (file_preview.MAX_TEXT + 10))
        self.assertEqual(len(data["content"]), file_preview.MAX_TEXT)
        self.assertTrue(data["truncated"])

    def test_utf16_is_supported(self):
        data = self.preview("notes.txt", "实验结果".encode("utf-16"))
        self.assertEqual(data["content"], "实验结果")

    def test_notebook_does_not_return_active_html_outputs(self):
        data = self.preview("analysis.ipynb", json.dumps({"cells": [{"cell_type": "code", "source": ["print(1)"], "outputs": [{"data": {"text/html": "<script>attack()</script>", "text/plain": ["1"]}}]}]}))
        self.assertEqual(data["kind"], "notebook")
        self.assertEqual(data["cells"][0]["output"], "1")

    def test_active_documents_are_not_inline_media(self):
        for extension in ["html", "svg", "js"]:
            self.assertNotIn("." + extension, file_preview.MEDIA)
            self.assertEqual(self.preview("untrusted." + extension, "<script>attack()</script>")["kind"], "text")

    def test_unknown_binary_is_download_only(self):
        self.assertEqual(self.preview("model.bin", b"\x00\x01\x02")["kind"], "download")

    def test_docx_paragraphs_and_tables_preserve_document_order(self):
        from docx import Document
        doc = Document()
        doc.add_heading("Results", 1)
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "latency"
        doc.add_paragraph("Conclusion")
        doc.save(os.path.join(self.temp.name, "paper.docx"))
        data = self.preview("paper.docx")
        self.assertEqual([block["type"] for block in data["blocks"]], ["paragraph", "table", "paragraph"])
        self.assertTrue(data["blocks"][0]["heading"])

    def test_excel_supports_multiple_sheets_without_evaluating_formulas(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(["node", "duration"])
        book.active.append(["node104", 1.25])
        second = book.create_sheet("Summary")
        second.append(["=1+1"])
        book.save(os.path.join(self.temp.name, "results.xlsx"))
        data = self.preview("results.xlsx")
        self.assertEqual(len(data["sheets"]), 2)
        self.assertEqual(data["sheets"][0]["rows"][1], ["node104", "1.25"])
        self.assertEqual(data["sheets"][1]["rows"][0], [""])

    def test_powerpoint_content(self):
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Experiment results"
        deck.save(os.path.join(self.temp.name, "report.pptx"))
        data = self.preview("report.pptx")
        self.assertEqual(data["kind"], "slides")
        self.assertIn("Experiment results", data["slides"][0])

    def test_archive_expansion_limit(self):
        path = os.path.join(self.temp.name, "large.docx")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("bomb.txt", b"0" * (65 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "超过"):
            self.preview("large.docx")
