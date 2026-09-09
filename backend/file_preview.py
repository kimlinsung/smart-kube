"""Bounded, read-only previews. Never execute uploaded code or Office macros."""
from __future__ import annotations

import csv
import io
import json
import os
import zipfile
from itertools import islice

MAX_TEXT = 1024 * 1024
MEDIA = {".pdf": ("pdf", "application/pdf"), ".png": ("image", "image/png"),
         ".jpg": ("image", "image/jpeg"), ".jpeg": ("image", "image/jpeg"),
         ".gif": ("image", "image/gif"), ".webp": ("image", "image/webp"),
         ".avif": ("image", "image/avif"), ".bmp": ("image", "image/bmp"),
         ".mp4": ("video", "video/mp4"), ".webm": ("video", "video/webm"),
         ".mp3": ("audio", "audio/mpeg"), ".wav": ("audio", "audio/wav"),
         ".ogg": ("audio", "audio/ogg")}


def _check_archive(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > 64 * 1024 * 1024:
            raise ValueError("文档解压后超过预览限制，请下载查看")
        if any(item.flag_bits & 1 for item in entries):
            raise ValueError("加密文档需要解密后上传")


def read_preview(item):
    path, name = item["stored_path"], item["original_name"]
    ext = os.path.splitext(name)[1].lower()
    result = {"filename": name, "size": os.path.getsize(path), "extension": ext, "truncated": False}
    remaining = MAX_TEXT

    def take(value, limit=2000):
        nonlocal remaining
        text = str(value) if value is not None else ""
        cropped = text[:min(limit, remaining)]
        remaining -= len(cropped)
        result["truncated"] |= len(cropped) < len(text)
        return cropped
    if result["size"] > 20 * 1024 * 1024:
        return {**result, "kind": "download", "content": "文件超过 20 MB，请下载查看"}
    if ext in MEDIA:
        return {**result, "kind": MEDIA[ext][0]}
    if ext in {".docx", ".xlsx", ".pptx"}:
        _check_archive(path)
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        blocks, chars = [], 0
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for element in doc.element.body:
            if len(blocks) >= 300 or chars >= MAX_TEXT:
                result["truncated"] = True
                break
            if element.tag.endswith("}p"):
                paragraph = Paragraph(element, doc)
                text = take(paragraph.text, 20000)
                blocks.append({"type": "paragraph", "text": text,
                               "heading": paragraph.style.name.startswith("Heading") if paragraph.style else False})
                chars += len(text)
            elif element.tag.endswith("}tbl"):
                table = Table(element, doc)
                rows = [[take(cell.text) for cell in row.cells[:30]] for row in table.rows[:100]]
                result["truncated"] |= len(table.rows) > 100 or any(len(row.cells) > 30 for row in table.rows[:100])
                blocks.append({"type": "table", "rows": rows})
                chars += sum(len(cell) for row in rows for cell in row)
        return {**result, "kind": "document", "blocks": blocks}
    if ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheets = []
        try:
            for sheet in book.worksheets[:10]:
                if not remaining:
                    result["truncated"] = True
                    break
                rows = [[take(value) for value in row]
                        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 201, 201), max_col=min(sheet.max_column or 40, 40), values_only=True)]
                sheets.append({"name": sheet.title, "rows": rows[:200],
                               "truncated": len(rows) > 200 or (sheet.max_column or 0) > 40})
            result["truncated"] |= len(book.worksheets) > 10 or any(s["truncated"] for s in sheets)
        finally:
            book.close()
        return {**result, "kind": "spreadsheet", "sheets": sheets}
    if ext == ".pptx":
        from pptx import Presentation
        deck = Presentation(path)
        slides = []
        for slide in list(deck.slides)[:40]:
            blocks = []
            for shape in list(slide.shapes)[:60]:
                if shape.has_text_frame:
                    blocks.append(take(shape.text, 20000))
                elif shape.has_table:
                    blocks.append("\n".join(" | ".join(take(cell.text) for cell in list(row.cells)[:40]) for row in list(shape.table.rows)[:40]))
            slides.append(blocks)
        return {**result, "kind": "slides", "slides": slides, "truncated": result["truncated"] or len(deck.slides) > 40}
    with open(path, "rb") as handle:
        raw = handle.read(MAX_TEXT + 1)
    result["truncated"] = len(raw) > MAX_TEXT
    raw = raw[:MAX_TEXT]
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        content = raw.decode("utf-16", errors="replace")
    elif b"\0" in raw:
        return {**result, "kind": "download", "content": "此二进制格式暂不支持预览，请下载查看"}
    else:
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("gb18030", errors="replace")
    result.update(kind="text", content=content)
    if ext in {".md", ".markdown"}:
        result["kind"] = "markdown"
    elif ext in {".csv", ".tsv"}:
        rows = list(islice(csv.reader(io.StringIO(content), delimiter="\t" if ext == ".tsv" else ","), 201))
        result.update(kind="spreadsheet", sheets=[{"name": name, "rows": [[cell[:2000] for cell in row[:40]] for row in rows[:200]]}],
                      truncated=result["truncated"] or len(rows) > 200 or any(len(row) > 40 for row in rows))
    elif ext in {".json", ".ipynb"} and not result["truncated"]:
        try:
            value = json.loads(content)
            if ext == ".ipynb" and isinstance(value, dict) and isinstance(value.get("cells"), list):
                cells = []
                for cell in value["cells"][:100]:
                    if not isinstance(cell, dict):
                        continue
                    source = cell.get("source", "")
                    source = "".join(source) if isinstance(source, list) else str(source)
                    outputs = []
                    for output in cell.get("outputs", [])[:20]:
                        text = output.get("text", output.get("data", {}).get("text/plain", ""))
                        outputs.append("".join(text) if isinstance(text, list) else str(text))
                    cells.append({"type": cell.get("cell_type"), "source": source[:20000], "output": "\n".join(outputs)[:20000]})
                result.update(kind="notebook", cells=cells, truncated=len(value["cells"]) > 100)
            else:
                result.update(kind="json", content=json.dumps(value, indent=2, ensure_ascii=False))
        except (ValueError, TypeError, AttributeError):
            pass
    return result
